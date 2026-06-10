import socket
import logging
import hashlib
import os
from celery import shared_task
from django.conf import settings
from auditlog.models import AuditEvent
from .models import EvidenceVersion

logger = logging.getLogger(__name__)

def calculate_sha256(file_field):
    """
    Computes SHA-256 hash for the uploaded file stream.
    """
    hash_sha256 = hashlib.sha256()
    file_field.seek(0)
    for chunk in file_field.chunks():
        hash_sha256.update(chunk)
    return hash_sha256.hexdigest()

def scan_file_stream(file_stream):
    """
    Connects to ClamAV daemon via TCP socket and streams data using INSTREAM protocol.
    Returns (status, virus_name)
    """
    host = getattr(settings, 'CLAMAV_HOST', 'localhost')
    port = int(getattr(settings, 'CLAMAV_PORT', 3310))
    
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
        s.settimeout(30)
        s.connect((host, port))
    except Exception as e:
        logger.error(f"Failed to connect to ClamAV at {host}:{port}: {e}")
        return 'Error', f"ClamAV connection failed: {e}"
        
    try:
        # Request INSTREAM scan
        s.sendall(b'zINSTREAM\x00')
        
        # Stream file in chunks
        file_stream.seek(0)
        while True:
            chunk = file_stream.read(8192)
            if not chunk:
                break
            # Send chunk length (4-byte big-endian int) followed by chunk data
            s.sendall(len(chunk).to_bytes(4, 'big'))
            s.sendall(chunk)
            
        # Send empty chunk (4 zero bytes) to signify EOF
        s.sendall((0).to_bytes(4, 'big'))
        
        # Await response
        response = b''
        while True:
            data = s.recv(1024)
            if not data:
                break
            response += data
            
        response_str = response.decode('utf-8').strip()
        logger.info(f"ClamAV scan response: {response_str}")
        
        if 'OK' in response_str:
            return 'Clean', 'No virus found.'
        elif 'FOUND' in response_str:
            # stream: Eicar-Signature FOUND -> Eicar-Signature
            parts = response_str.split(':')
            virus_name = "Virus detected"
            if len(parts) > 1:
                virus_name = parts[1].replace('FOUND', '').strip()
            return 'Infected', virus_name
        else:
            return 'Error', f"Unexpected response: {response_str}"
            
    except Exception as e:
        logger.error(f"Error streaming file to ClamAV: {e}")
        return 'Error', f"Scan error: {e}"
    finally:
        s.close()


@shared_task
def scan_file_clamav(evidence_version_id):
    """
    Celery task to scan an uploaded file version and trigger extraction if clean.
    """
    try:
        version = EvidenceVersion.objects.get(id=evidence_version_id)
    except EvidenceVersion.DoesNotExist:
        return f"EvidenceVersion {evidence_version_id} not found."
        
    # Generate SHA-256 hash checksum
    if not version.sha256_hash:
        version.sha256_hash = calculate_sha256(version.file)
        version.save(update_fields=['sha256_hash'])
        
    # Perform ClamAV socket stream scan
    with version.file.open('rb') as f:
        status, results = scan_file_stream(f)
        
    version.status = status
    version.scan_results = results
    version.save(update_fields=['status', 'scan_results'])
    
    # Audit logging
    tenant = version.document.tenant
    user = version.uploaded_by
    
    # Create Audit Event record
    AuditEvent.objects.create(
        tenant=tenant,
        user=user,
        event_type='FILE_SCAN',
        action='SCAN_COMPLETE' if status == 'Clean' else 'MALWARE_ALERT',
        ip_address='127.0.0.1',
        payload={
            'document_id': version.document.id,
            'version_id': version.id,
            'file_name': version.file_name,
            'sha256_hash': version.sha256_hash,
            'status': status,
            'results': results
        }
    )
    
    if status == 'Clean':
        # Trigger text extraction task
        extract_text_task.delay(version.id)
        
    return f"Scan complete for version {version.id}: {status}"


@shared_task
def extract_text_task(evidence_version_id):
    """
    Celery task to extract text contents from PDF, Word, and Excel files.
    """
    try:
        version = EvidenceVersion.objects.get(id=evidence_version_id)
    except EvidenceVersion.DoesNotExist:
        return f"EvidenceVersion {evidence_version_id} not found."
        
    if version.status != 'Clean':
        return f"Version {version.id} is not Clean (current status: {version.status})."
        
    ext = os.path.splitext(version.file_name)[1].lower()
    text = ""
    
    try:
        file_path = version.file.path
        
        if ext == '.pdf':
            from pypdf import PdfReader
            reader = PdfReader(file_path)
            pages_text = []
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    pages_text.append(page_text)
            text = "\n".join(pages_text)
            
        elif ext in ['.docx', '.doc']:
            from docx import Document
            doc = Document(file_path)
            paragraphs_text = [p.text for p in doc.paragraphs]
            text = "\n".join(paragraphs_text)
            
        elif ext in ['.xlsx', '.xls']:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            cells_text = []
            for sheet in wb.worksheets:
                cells_text.append(f"--- Sheet: {sheet.title} ---")
                for row in sheet.iter_rows(values_only=True):
                    row_str = " ".join([str(val) for val in row if val is not None])
                    if row_str.strip():
                        cells_text.append(row_str)
            text = "\n".join(cells_text)
            
        elif ext in ['.txt', '.csv']:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                text = f.read()
                
        if text.strip():
            version.extracted_text = text
            version.save(update_fields=['extracted_text'])
            return f"Extracted {len(text)} characters from {version.file_name}."
            
    except Exception as e:
        logger.error(f"Error extracting text from {version.file_name}: {e}")
        return f"Extraction failed: {e}"
        
    return f"No text content extracted from {version.file_name}."
